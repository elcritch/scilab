#!/usr/bin/env python3


import shutil, re, sys, os, itertools, argparse, json, collections
from pathlib import Path

import openpyxl
from openpyxl import load_workbook

import scilab.tools.scriptrunner as ScriptRunner
# from scilab.tools.scriptrunner import RESEARCH, RAWDATA, debug

import logging

from scilab.tools.excel import *
from scilab.tools.project import *

import scilab.tools.jsonutils as Json
from scilab.tools.tables import mdBlock, mdHeader, ImageTable, MarkdownTable
    
## Main
        
def parse_data_from_worksheet(parser_data_sheet_excel, testconf, args, **kwargs):

    try:
        excelfile = testconf.folder.datasheet 
        excelfile = excelfile.resolve()
        debug(excelfile)
        
        wb = load_workbook(excelfile.absolute().as_posix(), data_only=True)
    except (Exception) as err:
        logging.warn("Cannot open file:\n\t"+str(excelfile)+"\n\t vs \n\t"+str(excelfile))
        return 
        
    ## Process Excel Sheets
    ws = wb.worksheets[0]
    return parser_data_sheet_excel(ws)    



def parse_from_image_measurements(parser_image_measurements, testconf, args):

    # imgMeasureFile = testconf.folder.image / 'processed' / (testconf.info.name + '.measurements.json')
    imgMeasureFile = testconf.folder.images / 'processed' / 'data.json'
    imgMeasureFile = imgMeasureFile.resolve()        
    
    imgMeasurements = Json.load_json(imgMeasureFile.parent.as_posix(), json_url=imgMeasureFile.name)
    # debug(imgMeasurements)
    
    data = parser_image_measurements(testconf, imgMeasurements)

    return data
    
def HTML(arg):
    return arg.replace('\n','')

def handler(testconf, excelfile, args):
        
    def updateMetaData(data):
        ## Handle Names    
        data['name'] = testconf.info.name
        data['id'] = testconf.info.short
        data['info'] = testconf.info.as_dict()
        

    ## Update with Excel Values    
    data = parse_data_from_worksheet(
        parser_data_sheet_excel=args.parser_data_sheet_excel,
        testpath=excelfile,
        testconf=testconf,
        args=args,
        )

    print(mdBlock("Excel Sheet Data:"))
    rows = [ (k,v) for k,v in flatten(data).items() ]
    rows = sorted(rows)

    print()
    print(HTML(tabulate.tabulate( rows, [ "Key", "Value", ], tablefmt ='html' )))
    
    # print(mdBlock("<pre>\n"+json.dumps(data,indent=4)+"\n</pre>"))
    updateMetaData(data)
    
    # excel_data = DataTree(notes=data.pop("notes"), other=data.pop("other"))
    json_data = collections.OrderedDict(sorted(data.items()))
    testconf.folder.save_calculated_json_raw(test=testconf, name='excel',json_data=json_data, overwrite=True)
    
    ## Update with Image Measurements
    
    data = parse_from_image_measurements(
        parser_image_measurements=args.parser_image_measurements,
        testconf=testconf,
        args=args)
    
    updateMetaData(data)

    testconf.folder.save_calculated_json_raw(test=testconf, name='measurements',json_data=data, overwrite=True)
    
    return

    
